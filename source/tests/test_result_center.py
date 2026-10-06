"""Phase 6 projections, event-driven UI, account-safe menus and exact report schemas."""
import csv
from types import SimpleNamespace
from unittest import mock

import pytest

import client_app as module
import result_center as results
from test_account_management_ui import management_app, raw_account
from test_account_normalization import AccountTree


def center_app(tmp_path=None):
    app = management_app()
    if tmp_path:
        app.account_states.history = module.AccountHistoryStore(tmp_path / "history.db")
    app.run_config = module.RunConfig({"page_target": 3, "parsed_accounts": {
        index: app.account_states.get(index) for index in app.account_states.indexes()}})
    app.create_page_results_window = mock.Mock()
    app.create_page_results_window.winfo_exists.return_value = True
    app.create_page_result_trees = {key: AccountTree() for key, _ in results.RESULT_TABS}
    for tree in app.create_page_result_trees.values():
        tree.identify_row = mock.Mock()
        tree.selection_set = mock.Mock()
    app.result_center_summary_labels = {key: mock.Mock() for key in results.SUMMARY_KEYS}
    app.create_page_results_notebook = mock.Mock()
    app.create_page_results_notebook.index.return_value = 0
    app.lbl_create_page_result_summary = mock.Mock()
    app.result_center_menu = mock.Mock()
    app.refresh_create_page_results_dialog()
    return app


def apply_ui(app):
    app.post_ui.call_args.args[0]()


def values(app, index, tab="overview"):
    return app.create_page_result_trees[tab].item(app.result_center_item_id(app.account_states.get(index)), "values")


def page_outcome(page_id="99999", status="SUCCESS", reason=""):
    return {"status": status, "page_name": "Page A", "page_id": page_id,
            "page_url": "https://www.facebook.com/profile.php?id=" + page_id if page_id else "",
            "reason": reason, "page_job_id": "job-A", "verified": status == "SUCCESS"}


def test_submit_and_verified_success_update_without_waiting_for_batch():
    app = center_app()
    app.set_account_state(1, status="LIVE", current_action="Create Page [SUBMIT]")
    apply_ui(app)
    app.record_task_result(1, "CREATE_PAGE", "RUNNING", "submitted")
    apply_ui(app)
    a, b = values(app, 1), values(app, 2)
    assert a[2] == "LIVE" and a[8] == "RUNNING"
    assert a[4:8] == (0, 0, 0, 0) and a[11] == ""
    assert a[9] == "Create Page [SUBMIT]"
    assert b[2] == "UNKNOWN" and b[4:8] == (0, 0, 0, 0)
    app.result_center_summary_labels["RUNNING"].config.assert_called_with(text="1")
    app.record_task_result(1, "CREATE_PAGE", "SUCCESS", "verified", page_outcome())
    apply_ui(app)
    a = values(app, 1, "create_page")
    assert a[4:6] == (1, 0)
    assert a[8] == "SUCCESS" and a[10:12] == ("Page A", "99999")
    assert values(app, 2) == b
    app.result_center_summary_labels["PAGE SUCCESS"].config.assert_called_with(text="1")
    app.result_center_summary_labels["RUNNING"].config.assert_called_with(text="0")


def test_policy_rejection_is_failed_task_not_fake_success_or_die():
    app = center_app()
    app.set_account_state(1, status="LIVE")
    apply_ui(app)
    app.record_task_result(1, "CREATE_PAGE", "FAILED", "POLICY_REJECTED",
                           page_outcome("", "FAILED", "POLICY_REJECTED"))
    apply_ui(app)
    row = values(app, 1)
    assert row[2] == "LIVE"
    assert row[4:6] == (0, 1) and row[8] == "FAILED"
    assert row[11] == "" and row[12] == "POLICY_REJECTED"
    tree = app.create_page_result_trees["overview"]
    assert tree.item(app.result_center_item_id(app.account_states.get(1)), "tags") == ("TASK_FAILED",)
    assert not app.create_page_result_trees["error"].get_children()


@pytest.mark.parametrize("status,tab", [("CHECKPOINT", "checkpoint"), ("DIE", "error"), ("ERROR", "error")])
def test_terminal_account_events_update_filter_and_reason(status, tab):
    app = center_app()
    app.set_account_state(1, status=status, current_action="account-reason-A")
    apply_ui(app)
    assert values(app, 1, tab)[2] == status
    assert values(app, 1, tab)[12] == "account-reason-A"
    assert values(app, 2)[2] == "UNKNOWN"
    app.result_center_summary_labels[status].config.assert_called_with(text="1")
    app.result_center_summary_labels["TOTAL ACCOUNTS"].config.assert_called_with(text="2")


def test_friend_and_page_counters_are_account_local_and_realtime():
    app = center_app()
    app.record_task_result(1, "CREATE_PAGE", "SUCCESS", "verified", page_outcome())
    apply_ui(app)
    app.record_task_result(2, "FRIEND_REQUEST", "SUCCESS", "sent", {"target": "33333", "result": "SENT"})
    apply_ui(app)
    app.record_task_result(2, "FRIEND_REQUEST", "FAILED", "not verified", {"target": "44444"})
    apply_ui(app)
    assert values(app, 1)[4:8] == (1, 0, 0, 0)
    assert values(app, 2, "friend")[4:8] == (0, 0, 1, 1)
    assert len(app.create_page_result_trees["create_page"].get_children()) == 1
    assert len(app.create_page_result_trees["friend"].get_children()) == 1


def test_incremental_update_keeps_other_rows_and_stable_ids_after_reorder():
    app = center_app()
    tree = app.create_page_result_trees["overview"]
    before = dict(tree.rows)
    original_item = app.result_center_item_id(app.account_states.get(1))
    app.set_account_state(1, status="CHECKING", current_action="checking A")
    apply_ui(app)
    b_item = app.result_center_item_id(app.account_states.get(2))
    assert tree.rows[b_item] is before[b_item]
    app.account_states.sync([dict(app.account_states.get(2), stt=1), dict(app.account_states.get(1), stt=2)])
    app.refresh_create_page_results_dialog()
    assert tree.item(original_item, "values")[:3] == (2, "11111", "CHECKING")
    assert len(tree.get_children()) == 2


def test_removed_processed_input_rows_are_read_from_existing_history(tmp_path):
    app = center_app(tmp_path)
    app.account_states.set_login_mode(1, "COOKIE")
    app.set_account_state(1, status="LIVE")
    app.record_task_result(1, "CREATE_PAGE", "SUCCESS", "verified", page_outcome())
    app.account_states.sync([])
    app.refresh_create_page_results_dialog()
    state = next(state for state in app.result_center_accounts() if state["uid"] == "11111")
    assert state["page_success_count"] == 1 and state["status"] == "LIVE"
    assert state["login_mode"] == "COOKIE"
    assert state["tasks"]["CREATE_PAGE"]["outcomes"][0]["page_id"] == "99999"
    assert app.account_states.indexes() == []  # Projection never restores/mutates runtime state.
    app.result_center_summary_labels["PAGE SUCCESS"].config.assert_called_with(text="1")


@pytest.mark.parametrize("mode,expected", [
    ("uid", "11111"), ("canonical", raw_account("11111", "A")),
    ("page_id", "99999"), ("page_url", "https://www.facebook.com/profile.php?id=99999"),
    ("reason", ""),
])
def test_right_click_copy_is_exact_clicked_account_even_after_reorder(mode, expected):
    app = center_app()
    app.record_task_result(1, "CREATE_PAGE", "SUCCESS", "verified", page_outcome())
    apply_ui(app)
    tree = app.create_page_result_trees["overview"]
    item = app.result_center_item_id(app.account_states.get(1))
    tree.identify_row.return_value = item
    app.show_result_center_context_menu(SimpleNamespace(y=12, x_root=10, y_root=20), tree)
    tree.selection_set.assert_called_once_with(item)
    app.result_center_menu.grab_release.assert_called_once()
    app.account_states.sync([dict(app.account_states.get(2), stt=1), dict(app.account_states.get(1), stt=2)])
    app.copy_result_center_data(mode)
    app.root.clipboard_append.assert_called_once_with(expected)


def test_copy_reason_and_full_result_row_have_no_credentials():
    app = center_app()
    app.record_task_result(1, "CREATE_PAGE", "FAILED", "password-A token-A policy rejected")
    apply_ui(app)
    tree = app.create_page_result_trees["overview"]
    item = app.result_center_item_id(app.account_states.get(1))
    tree.identify_row.return_value = item
    app.show_result_center_context_menu(SimpleNamespace(y=12, x_root=10, y_root=20), tree)
    app.copy_result_center_data("reason")
    assert app.root.clipboard_append.call_args.args[0] == "[REDACTED] [REDACTED] policy rejected"
    app.copy_result_center_data("row")
    assert app.root.clipboard_append.call_args.args[0] == "\t".join(map(str, tree.item(item, "values")))
    assert "password-A" not in str(app.root.clipboard_append.call_args)
    assert "token-A" not in str(app.root.clipboard_append.call_args)


def test_stale_or_empty_result_row_never_copies_other_account():
    app = center_app()
    tree = app.create_page_result_trees["overview"]
    item = app.result_center_item_id(app.account_states.get(1))
    row = list(tree.item(item, "values"))
    row[1] = "22222"
    tree.item(item, values=row)
    tree.identify_row.return_value = item
    app.show_result_center_context_menu(SimpleNamespace(y=12, x_root=10, y_root=20), tree)
    app.copy_result_center_data("uid")
    app.root.clipboard_append.assert_not_called()
    app.result_center_menu.tk_popup.assert_not_called()


def test_result_export_schema_status_is_d_not_password_column(tmp_path):
    app = center_app()
    state = app.account_states.get(1)
    state.update(status="CHECKPOINT", password="=secret", raw_line="=raw", last_error="checkpoint")
    path = tmp_path / "results.xlsx"
    rows = [results.result_values(state)]
    module.write_create_page_results_xlsx(path, {"checkpoint": [results.report_record(state, 3)]}, rows)
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook["CHECKPOINT"]
        assert tuple(cell.value for cell in sheet[1]) == results.REPORT_HEADERS
        assert sheet["A2"].value == 1 and sheet["B2"].value == "11111"
        assert sheet["C2"].value == "=secret" and sheet["C2"].data_type == "s"
        assert sheet["D2"].value == "CHECKPOINT"
        assert sheet["F2"].value == 3 and sheet["G2"].value == "checkpoint"
        assert sheet["I2"].value == "=raw" and sheet["I2"].data_type == "s"
        detail = workbook["RESULTS"]
        assert tuple(cell.value for cell in detail[1]) == results.RESULT_HEADERS
        assert tuple(cell.value if cell.value is not None else "" for cell in detail[2]) == rows[0]
    finally:
        workbook.close()


@pytest.mark.parametrize("failure", [None, "save", "append", "empty"])
def test_workbook_always_closes_including_failure(tmp_path, failure):
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.close = mock.Mock(wraps=workbook.close)
    if failure == "save":
        workbook.save = mock.Mock(side_effect=OSError("save failed"))
    if failure == "append":
        workbook.create_sheet = mock.Mock(side_effect=ValueError("append failed"))
    records = {} if failure == "empty" else {"die": [{"uid": "11111", "reason": "expired"}]}
    with mock.patch.object(openpyxl, "Workbook", return_value=workbook):
        if failure:
            with pytest.raises((OSError, ValueError)):
                module.write_create_page_results_xlsx(tmp_path / "test.xlsx", records)
        else:
            module.write_create_page_results_xlsx(tmp_path / "test.xlsx", records)
    workbook.close.assert_called_once()


def test_csv_matches_all_fourteen_result_fields(tmp_path):
    app = center_app()
    rows = [results.result_values(app.account_states.get(index)) for index in (1, 2)]
    path = tmp_path / "results.csv"
    results.write_result_csv(path, rows)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream))
    assert tuple(records[0]) == results.RESULT_HEADERS
    assert records[1:] == [[str(value) for value in row] for row in rows]


@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_export_uses_selected_filter_and_one_consistent_snapshot(extension, tmp_path):
    app = center_app()
    app.set_account_state(1, status="CHECKPOINT", current_action="reason A")
    app.create_page_results_notebook.index.return_value = 2
    path = tmp_path / ("filtered." + extension)
    def file_dialog(**_kwargs):
        app.account_states.update(2, status="ERROR", current_action="changed during dialog")
        return str(path)
    with mock.patch.object(module.filedialog, "asksaveasfilename", side_effect=file_dialog), \
         mock.patch.object(module.messagebox, "showinfo"), \
         mock.patch.object(module, "write_create_page_results_xlsx") as xlsx, \
         mock.patch.object(module, "write_result_csv") as csv_writer:
        app.export_result_center(extension)
    if extension == "xlsx":
        records = xlsx.call_args.args[1]["checkpoint"]
        assert len(records) == 1 and records[0]["uid"] == "11111"
        assert records[0]["status"] == "CHECKPOINT" and records[0]["target_count"] == 3
        assert xlsx.call_args.kwargs["result_rows"][0][12] == "reason A"
        csv_writer.assert_not_called()
    else:
        rows = csv_writer.call_args.args[1]
        assert len(rows) == 1 and rows[0][1] == "11111" and rows[0][2] == "CHECKPOINT"
        xlsx.assert_not_called()


def test_result_center_real_tk_schema_palette_reopen_and_menu(tmp_path):
    root = module.tk.Tk()
    root.withdraw()
    real_top = module.tk.Toplevel
    def hidden_top(parent):
        window = real_top(parent)
        window.withdraw()
        return window
    try:
        with mock.patch.object(module, "APP_DATA_DIR", str(tmp_path)), \
             mock.patch.object(module.MainToolApp, "load_settings"), \
             mock.patch.object(module, "check_for_updates"), \
             mock.patch.object(module.threading, "Thread"), \
             mock.patch.object(module.tk, "Toplevel", side_effect=hidden_top):
            app = module.MainToolApp(root, "2027-01-01")
            app.open_create_page_results_dialog()
            window = app.create_page_results_window
            app.open_create_page_results_dialog()
            assert app.create_page_results_window is window
            assert tuple(app.create_page_result_trees) == tuple(key for key, _ in results.RESULT_TABS)
            for tree in app.create_page_result_trees.values():
                assert tuple(tree.cget("columns")) == results.RESULT_COLUMNS
                assert tuple(tree.heading(column, "text") for column in results.RESULT_COLUMNS) == results.RESULT_HEADERS
                assert str(tree.tag_configure("LIVE", "background")) == module.ACCOUNT_STATUS_PALETTE["LIVE"]["background"]
                assert str(tree.tag_configure("TASK_FAILED", "background")) == module.ACCOUNT_STATUS_PALETTE["ERROR"]["background"]
                assert tree.cget("xscrollcommand") and tree.cget("yscrollcommand")
            labels = [app.result_center_menu.entrycget(index, "label") for index in range(app.result_center_menu.index("end") + 1)
                      if app.result_center_menu.type(index) == "command"]
            assert labels == ["Copy UID", "Copy canonical account", "Copy Page ID", "Copy Page URL", "Copy reason",
                              "Copy result row", "Mở Chrome kiểm tra", "Xem log account"]
            app._result_center_menu_account = {"uid": "11111"}
            with mock.patch.object(app, "open_account_inspector") as inspect, \
                 mock.patch.object(app, "show_management_account_logs") as logs:
                for index in range(app.result_center_menu.index("end") + 1):
                    if app.result_center_menu.type(index) == "command" and app.result_center_menu.entrycget(index, "label") in {"Mở Chrome kiểm tra", "Xem log account"}:
                        app.result_center_menu.invoke(index)
                inspect.assert_called_once_with(app._result_center_menu_account)
                logs.assert_called_once_with(app._result_center_menu_account)
            root.update_idletasks()
            app.close_requested = True
    finally:
        root.destroy()
