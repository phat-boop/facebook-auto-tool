"""Read-only result projections and report writers; never verify automation success."""
import csv
from datetime import datetime

from account_history import COUNTER_FIELDS, redact_history_text


RESULT_COLUMNS = (
    "stt", "uid", "status", "login_mode", "page_success", "page_failed",
    "friend_success", "friend_failed", "result", "action", "page_name", "page_id",
    "reason", "last_update",
)
RESULT_HEADERS = (
    "STT", "UID", "STATUS", "LOGIN MODE", "PAGE SUCCESS", "PAGE FAILED",
    "FRIEND SUCCESS", "FRIEND FAILED", "RESULT", "ACTION", "LAST PAGE NAME",
    "LAST PAGE ID", "REASON", "LAST UPDATE",
)
RESULT_TABS = (
    ("overview", "TỔNG QUAN"), ("create_page", "CREATE PAGE"),
    ("checkpoint", "CHECKPOINT"), ("error", "DIE / ERROR"), ("friend", "FRIEND"),
)
REPORT_HEADERS = (
    "STT", "UID", "Mật khẩu (Pass)", "Trạng thái", "Page đã tạo",
    "Chỉ tiêu Page", "Kết quả / Lý do", "Thời gian", "Dữ liệu gốc",
)
SUMMARY_KEYS = ("TOTAL ACCOUNTS", "RUNNING", "LIVE", "DIE", "CHECKPOINT", "ERROR", "PAGE SUCCESS", "PAGE FAILED")


def result_values(state):
    tasks = state.get("tasks", {})
    page_task = tasks.get("CREATE_PAGE", {})
    last_page = next((value for value in reversed(page_task.get("outcomes", []))
                      if isinstance(value, dict)), {})
    result = state.get("last_task_result", "PENDING")
    if state.get("status") in {"DIE", "ERROR", "CHECKPOINT"}:
        reason = state.get("last_error") or state.get("current_action")
    else:
        failed_task = next((task for task in tasks.values() if task.get("status") == "ERROR"), None)
        failed_task = failed_task or next((task for task in tasks.values() if task.get("status") == "FAILED"), {})
        reason = failed_task.get("detail") or last_page.get("reason") or last_page.get("technical_error") or ""
    timestamp = str(state.get("last_update") or "")
    try:
        timestamp = datetime.fromisoformat(timestamp).astimezone().strftime("%Y-%m-%d %H:%M:%S") if timestamp else "-"
    except ValueError:
        timestamp = "-"
    return (
        state.get("stt", ""), state.get("uid") or state.get("account_id", ""),
        state.get("status", "UNKNOWN"), state.get("login_mode") or "-",
        *(int(state.get(field) or 0) for field in COUNTER_FIELDS), result,
        redact_history_text(state.get("current_action"), state),
        redact_history_text(last_page.get("page_name"), state), str(last_page.get("page_id") or ""),
        redact_history_text(reason, state), timestamp,
    )


def result_matches(state, tab):
    if tab == "overview":
        return True
    if tab == "checkpoint":
        return state.get("status") == "CHECKPOINT"
    if tab in {"error", "die"}:
        return state.get("status") in ({"DIE"} if tab == "die" else {"DIE", "ERROR"})
    if tab == "completed":
        return state.get("tasks", {}).get("CREATE_PAGE", {}).get("status") == "SUCCESS"
    prefix, module = {"create_page": ("page", "CREATE_PAGE"), "friend": ("friend", "FRIEND_REQUEST")}[tab]
    return module in state.get("tasks", {}) or any(state.get(f"{prefix}_{kind}_count", 0) for kind in ("success", "failed"))


def result_summary(accounts):
    summary = dict.fromkeys(SUMMARY_KEYS, 0)
    for state in accounts:
        summary["TOTAL ACCOUNTS"] += 1
        status = state.get("status", "UNKNOWN")
        if status in {"LIVE", "DIE", "CHECKPOINT", "ERROR"}:
            summary[status] += 1
        if status not in {"DIE", "ERROR", "CHECKPOINT"} and (
            status == "CHECKING" or any(task.get("status") == "RUNNING" for task in state.get("tasks", {}).values())
        ):
            summary["RUNNING"] += 1
        summary["PAGE SUCCESS"] += int(state.get("page_success_count") or 0)
        summary["PAGE FAILED"] += int(state.get("page_failed_count") or 0)
    return summary


def report_record(state, target_count=0):
    values = result_values(state)
    return {
        "stt": values[0], "uid": values[1], "password": state.get("password", ""),
        "status": values[2], "created_count": values[4], "target_count": target_count,
        "reason": values[12] or values[8], "time": values[13], "raw_line": state.get("raw_line", ""),
    }


def write_result_csv(file_path, rows):
    with open(file_path, "w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(RESULT_HEADERS)
        writer.writerows(rows)


def write_result_workbook(file_path, categorized_records, result_rows=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Thiếu thư viện openpyxl để xuất Excel.") from exc
    workbook = Workbook()
    try:
        workbook.remove(workbook.active)
        names = {"completed": "LIVE", "die": "DIE", "overview": "TONG_QUAN",
                 "create_page": "CREATE_PAGE", "checkpoint": "CHECKPOINT", "error": "DIE_ERROR", "friend": "FRIEND"}
        for key, records in categorized_records.items():
            sheet = workbook.create_sheet(names[key])
            sheet.append(REPORT_HEADERS)
            for record in records:
                sheet.append((
                    record.get("stt", ""), record.get("uid") or record.get("account_id", ""),
                    record.get("password") or record.get("pwd", ""), record.get("status") or names[key],
                    int(record.get("created_count") or 0), int(record.get("target_count") or 0),
                    record.get("reason") or ("Đã tạo đủ Page" if key == "completed" else "Tài khoản không còn đăng nhập hợp lệ"),
                    record.get("time", ""), record.get("raw_line", ""),
                ))
        if result_rows is not None:
            sheet = workbook.create_sheet("RESULTS")
            sheet.append(RESULT_HEADERS)
            for row in result_rows:
                sheet.append(tuple(row))
        if not workbook.sheetnames:
            raise ValueError("Không có danh sách để xuất Excel.")
        for sheet in workbook:
            widths = (8, 24, 24, 18, 16, 16, 48, 22, 58) if sheet.title != "RESULTS" else (
                8, 24, 18, 24, 20, 20, 22, 22, 18, 55, 32, 24, 55, 24,
            )
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="24323D")
                cell.font = Font(color="F4FAFC", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # User data must remain text, not Excel formulas.
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}"
            sheet.row_dimensions[1].height = 34
            for index, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(file_path)
    finally:
        workbook.close()
