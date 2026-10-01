import ast
import asyncio
import contextvars
import json
import pathlib
import tempfile
import threading
import unittest
from unittest import mock

import client_app


class CoreHelpersTest(unittest.TestCase):
    def test_version_comparison_is_numeric(self):
        self.assertGreater(client_app.version_tuple("2.10.0"), client_app.version_tuple("2.9.9"))

    def test_update_url_must_be_trusted_executable(self):
        self.assertTrue(
            client_app.is_trusted_update_url(
                "https://github.com/phat-boop/facebook-auto-tool/releases/download/v2.1.3/client_app.exe"
            )
        )
        self.assertFalse(client_app.is_trusted_update_url("http://github.com/file.exe"))
        self.assertFalse(client_app.is_trusted_update_url("https://example.com/file.exe"))
        self.assertFalse(client_app.is_trusted_update_url("https://github.com/releases"))

        self.assertIsNone(
            client_app.get_self_update_target(r"C:\Python\python.exe", frozen=False)
        )
        self.assertIsNone(
            client_app.get_self_update_target(r"C:\Python\python.exe", frozen=True)
        )
        packaged_target = r"C:\Apps\FacebookTool\client_app.exe"
        self.assertEqual(
            client_app.get_self_update_target(packaged_target, frozen=True),
            packaged_target,
        )
        restart_script = client_app.build_windows_update_restart_script(
            4321,
            r"C:\Users\Test User\AppData\Local\FacebookAutoTool\update_temp.exe",
            packaged_target,
            r"C:\Users\Test User\AppData\Local\FacebookAutoTool\update_restart_error.log",
        )
        copy_position = restart_script.index('copy /y ')
        copy_check_position = restart_script.index('if errorlevel 1 goto replace_failed')
        reset_position = restart_script.index('set "PYINSTALLER_RESET_ENVIRONMENT=1"')
        restart_position = restart_script.index(
            'start "" /d "C:\\Apps\\FacebookTool" "C:\\Apps\\FacebookTool\\client_app.exe"'
        )
        self.assertLess(copy_position, copy_check_position)
        self.assertLess(copy_check_position, reset_position)
        self.assertLess(reset_position, restart_position)
        self.assertIn('if errorlevel 1 goto restart_failed', restart_script)
        self.assertIn(':restart_failed', restart_script)
        self.assertIn('Vui long mo lai ung dung thu cong.', restart_script)
        self.assertIn('del /q "%~f0"', restart_script)

        class FakeResponse:
            status = 200

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def read(self):
                return json.dumps({"version": client_app.CURRENT_VERSION}).encode("utf-8")

        fake_app = mock.Mock()
        fake_app.post_ui.side_effect = lambda callback: callback()
        with mock.patch.object(client_app.urllib.request, "urlopen", return_value=FakeResponse()), \
             mock.patch.object(client_app.messagebox, "showinfo") as showinfo:
            client_app.check_for_updates(fake_app, manual=False)
            showinfo.assert_not_called()

            client_app.check_for_updates(fake_app, manual=True)
            showinfo.assert_called_once_with(
                "Kiểm tra cập nhật",
                f"Bạn đang sử dụng phiên bản mới nhất\n\nPhiên bản hiện tại: v{client_app.CURRENT_VERSION}",
                parent=fake_app.root,
            )

    def test_cookie_parser_preserves_values(self):
        cookies = client_app.parse_cookies("c_user=123; xs=abc=def;")
        self.assertEqual(cookies[0]["name"], "c_user")
        self.assertEqual(cookies[1]["value"], "abc=def")

        class FakeText:
            def __init__(self, value):
                self.value = value

            def get(self, *_args):
                return self.value

            def delete(self, *_args):
                raise AssertionError("raw account textbox must not be rewritten")

            def insert(self, _index, value):
                raise AssertionError(f"raw account textbox must not be rewritten: {value}")

        account_cases = [
            (
                "123456789|SecretPass|ABCDEFGHIJKLMNOP|||"
                "c_user=123456789; xs=abc=def;|EAAB_TOKEN_VALUE",
                {
                    "uid": "123456789",
                    "pwd": "SecretPass",
                    "2fa": "ABCDEFGHIJKLMNOP",
                    "cookie": "c_user=123456789; xs=abc=def;",
                    "token": "EAAB_TOKEN_VALUE",
                },
            ),
            (
                "123456789|SecretPass|ABCDEFGHIJKLMNOP|"
                "c_user=123456789; xs=abc=def;|mail@example.com|"
                "127.0.0.1:8080:user:pass",
                {
                    "uid": "123456789",
                    "pwd": "SecretPass",
                    "2fa": "ABCDEFGHIJKLMNOP",
                    "cookie": "c_user=123456789; xs=abc=def;",
                    "email": "mail@example.com",
                    "proxy": "127.0.0.1:8080:user:pass",
                },
            ),
            (
                "c_user=987654321; xs=cookie=value;",
                {"uid": "987654321", "cookie": "c_user=987654321; xs=cookie=value;"},
            ),
            (
                "FACEBOOK|raw@example.com|RawPassword!",
                {
                    "type": "RAW", "uid": "raw@example.com",
                    "pwd": "RawPassword!", "email": "raw@example.com",
                },
            ),
        ]

        for original, expected_fields in account_cases:
            with self.subTest(account=original.split("|", 1)[0]):
                fake_app = mock.Mock()
                fake_app.txt_accounts = FakeText(original)
                fake_app.reload_table_from_text = mock.Mock()

                client_app.MainToolApp.auto_format_cookie_numbers(fake_app)

                self.assertEqual(fake_app.txt_accounts.value, original)
                fake_app.reload_table_from_text.assert_called_once_with()
                before = client_app.parse_any_account_line(original, 1)
                serialized = client_app.serialize_account_line(before)
                after = client_app.parse_any_account_line(serialized, 1)
                self.assertEqual(serialized, original)
                self.assertEqual(after["raw_line"], original)
                self.assertEqual(after["fields"], original.split("|"))
                for field, expected_value in expected_fields.items():
                    self.assertEqual(after[field], expected_value, field)

        extended = client_app.parse_any_account_line(account_cases[0][0], 1)
        self.assertEqual(extended["fields"][3:5], ["", ""])
        self.assertEqual(len(extended["fields"]), 7)
        state_store = client_app.AccountStateStore()
        state_store.sync([{
            **extended,
            "stt": 1,
            "account_id": extended["uid"],
        }])
        state = state_store.get(1)
        for field in ("raw_line", "uid", "password", "2fa", "cookie", "token", "fields"):
            self.assertEqual(state[field], extended[field])

        multiple_raw = "\n".join(case[0] for case in account_cases)
        parsed_multiple = [
            client_app.parse_any_account_line(line, index)
            for index, line in enumerate(multiple_raw.splitlines(), 1)
        ]
        self.assertEqual(client_app.serialize_account_lines(parsed_multiple), multiple_raw)

        full_after = client_app.parse_any_account_line(
            "1. " + account_cases[1][0],
            1,
        )
        self.assertEqual(full_after["pwd"], "SecretPass")
        self.assertEqual(full_after["2fa"], "ABCDEFGHIJKLMNOP")
        self.assertEqual(full_after["email"], "mail@example.com")
        self.assertEqual(full_after["proxy"], "127.0.0.1:8080:user:pass")

        locale_cases = (
            ("vi-VN", "VN", "Asia/Ho_Chi_Minh"),
            ("en-US", "US", "America/New_York"),
            ("th-TH", "TH", "Asia/Bangkok"),
        )
        for locale, country, timezone in locale_cases:
            parsed = client_app.parse_any_account_line(
                f"123456789|pass|ABCDEFGHIJKLMNOP|c_user=123456789;|"
                f"mail@example.com|127.0.0.1:8080|locale={locale}|country={country}|"
                f"timezone={timezone}",
                1,
            )
            self.assertEqual(parsed["locale"], locale)
            self.assertEqual(parsed["country"], country)
            self.assertEqual(parsed["timezone"], timezone)
            self.assertEqual(parsed["proxy"], "127.0.0.1:8080")
        self.assertEqual(
            client_app.parse_any_account_line("c_user=123;", 1)["locale"],
            "AUTO",
        )
        self.assertIsNone(client_app.browser_locale_options("AUTO")["locale"])
        self.assertEqual(client_app.browser_locale_options("id-ID")["locale"], "id-ID")

    def test_proxy_parser_supports_authenticated_and_scheme_formats(self):
        self.assertEqual(
            client_app.parse_proxy("127.0.0.1:8080:user:pass"),
            {
                "server": "http://127.0.0.1:8080",
                "username": "user",
                "password": "pass",
            },
        )
        self.assertEqual(
            client_app.parse_proxy("socks5://user:pass@127.0.0.1:1080")["server"],
            "socks5://127.0.0.1:1080",
        )
        self.assertIsNone(client_app.parse_proxy("127.0.0.1:99999"))
        resolved = {1: "10.0.0.1:8001", 2: "10.0.0.2:8002"}
        self.assertEqual(client_app.resolve_account_proxy("", resolved, 1), resolved[1])
        self.assertEqual(client_app.resolve_account_proxy("", resolved, 2), resolved[2])
        self.assertNotEqual(
            client_app.resolve_account_proxy("", resolved, 1),
            client_app.resolve_account_proxy("", resolved, 2),
        )
        self.assertEqual(
            client_app.resolve_account_proxy("10.0.0.9:9000", resolved, 1),
            "10.0.0.9:9000",
        )

    def test_account_state_isolation_failures_and_batches(self):
        self.assertEqual(client_app.MainToolApp._bounded_int("100", 6, 1, 20), 20)
        self.assertEqual(
            client_app.effective_account_worker_count(20, {"create_page": True}, 3),
            3,
        )
        self.assertEqual(
            client_app.effective_account_worker_count(8, {"create_page": False}, 3),
            8,
        )

        states = client_app.AccountStateStore()
        states.sync([
            {
                "stt": 1, "account_id": "account-A", "locale": "vi-VN",
                "country": "VN", "proxy": "10.0.0.1:8001",
                "raw_line": "1. account-A|pass-A", "source_line": "account-A|pass-A",
            },
            {
                "stt": 2, "account_id": "account-B", "locale": "th-TH",
                "country": "TH", "proxy": "10.0.0.2:8002",
                "raw_line": "2. account-B|pass-B", "source_line": "account-B|pass-B",
            },
        ])
        states.append_log(1, "log chỉ thuộc A")
        self.assertIn("log chỉ thuộc A", states.get(1)["logs"])
        self.assertNotIn("log chỉ thuộc A", states.get(2)["logs"])

        states.update(1, status="LIVE", current_action="Đang chạy")
        self.assertEqual(states.get(1)["status"], "LIVE")
        self.assertEqual(states.get(2)["status"], "UNKNOWN")
        self.assertEqual(states.get(1)["locale"], "vi-VN")
        self.assertEqual(states.get(2)["locale"], "th-TH")
        self.assertEqual(states.get(1)["proxy"], "10.0.0.1:8001")
        self.assertEqual(states.get(2)["proxy"], "10.0.0.2:8002")
        self.assertEqual(states.get(1)["source_line"], "account-A|pass-A")
        self.assertEqual(states.get(1)["name"], "account-A")
        self.assertEqual(
            client_app.account_display_name({"account_id": "legacy-account"}, 9),
            "legacy-account",
        )
        self.assertEqual(client_app.account_display_name({}, 9), "Nick_9")

        self.assertEqual(
            client_app.keep_unprocessed_account_lines(
                ["1. account-A|pass-A", "2. account-B|pass-B", "3. account-C|pass-C"],
                {"account-A|pass-A", "account-B|pass-B"},
            ),
            ["3. account-C|pass-C"],
        )
        self.assertTrue(client_app.is_invalid_facebook_account_url(
            "https://www.facebook.com/disabled/"
        ))
        self.assertTrue(client_app.is_invalid_facebook_account_url(
            "https://www.facebook.com/checkpoint/"
        ))
        self.assertFalse(client_app.is_invalid_facebook_account_url(
            "https://www.facebook.com/pages/creation/"
        ))

        result_app = client_app.MainToolApp.__new__(client_app.MainToolApp)
        result_app.account_states = states
        result_app.create_page_result_lock = threading.RLock()
        result_app.create_page_account_results = {"completed": {}, "die": {}}
        result_app.post_ui = lambda _callback: None
        result_app.record_create_page_account_result(
            1, "completed", created_count=3, target_count=3
        )
        result_app.record_create_page_account_result(
            2, "die", reason="Cookie hết hạn", created_count=1, target_count=3
        )
        completed = result_app.get_create_page_account_results("completed")
        died = result_app.get_create_page_account_results("die")
        self.assertEqual(completed[0]["account_id"], "account-A")
        self.assertEqual(completed[0]["created_count"], 3)
        self.assertEqual(died[0]["account_id"], "account-B")
        self.assertEqual(died[0]["reason"], "Cookie hết hạn")
        self.assertEqual(
            result_app.get_processed_create_page_sources(),
            {"account-A|pass-A", "account-B|pass-B"},
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = pathlib.Path(temp_dir) / "create_page_results.xlsx"
            client_app.write_create_page_results_xlsx(
                report_path,
                {"completed": completed, "die": died},
            )
            from openpyxl import load_workbook
            workbook = load_workbook(report_path, read_only=True)
            self.assertEqual(workbook.sheetnames, ["LIVE", "DIE"])
            self.assertEqual(workbook["LIVE"]["B2"].value, "account-A")
            self.assertEqual(workbook["LIVE"]["C2"].value, "LIVE")
            self.assertEqual(workbook["DIE"]["B2"].value, "account-B")
            self.assertEqual(workbook["DIE"]["C2"].value, "DIE")
            workbook.close()

        async def update_scoped_account(index, status, message):
            await asyncio.sleep(0)
            states.update(index, status=status, current_action=message)
            states.append_log(index, message)

        async def update_both_accounts():
            await asyncio.gather(
                update_scoped_account(1, "LIVE", "vi account action"),
                update_scoped_account(2, "CHECKING", "th account action"),
            )

        asyncio.run(update_both_accounts())
        self.assertIn("vi account action", states.get(1)["logs"])
        self.assertNotIn("vi account action", states.get(2)["logs"])
        self.assertIn("th account action", states.get(2)["logs"])
        self.assertNotIn("th account action", states.get(1)["logs"])

        self.assertEqual(states.set_failure(1, "proxy", "Proxy timeout"), "ERROR")
        self.assertEqual(states.get(1)["status"], "ERROR")
        self.assertNotEqual(states.get(1)["status"], "DIE")
        self.assertEqual(states.set_failure(1, "invalid_cookie", "Cookie hết hạn"), "DIE")
        self.assertEqual(states.set_failure(2, "invalid_login", "Sai đăng nhập"), "DIE")

        batches = client_app.split_account_batches(range(1, 21), 5)
        self.assertEqual(len(batches), 4)
        self.assertTrue(all(len(batch) == 5 for batch in batches))

        class FakeLocator:
            def __init__(self, count=1, on_click=None):
                self._count = count
                self._on_click = on_click

            @property
            def first(self):
                return self

            async def fill(self, _value):
                return None

            async def count(self):
                return self._count

            async def click(self):
                if self._on_click:
                    self._on_click()

            async def is_visible(self):
                return True

        class FakePage:
            def __init__(self, final_url, login_error=None):
                self.url = "about:blank"
                self.final_url = final_url
                self.login_error = login_error

            async def goto(self, url, **_kwargs):
                if self.login_error:
                    raise self.login_error
                self.url = url

            def locator(self, selector):
                if selector == 'button[name="login"], button[type="submit"]':
                    return FakeLocator(on_click=lambda: setattr(self, "url", self.final_url))
                if selector == 'input[name="email"], input[name="pass"], form[action*="login"]':
                    return FakeLocator(count=0)
                return FakeLocator()

            async def wait_for_timeout(self, _milliseconds):
                return None

        class FakeContext:
            async def cookies(self, _url):
                return [{"name": "c_user", "value": "123456789"}]

            async def close(self):
                return None

        class FakeBrowser:
            async def close(self):
                return None

        async def run_login_case(final_url, login_error=None):
            app = client_app.MainToolApp.__new__(client_app.MainToolApp)
            app.is_running = True
            app.stop_requested = False
            app.run_config = {
                "modes": {}, "options": {}, "targets": "", "headless": True,
                "target": 1, "min_delay": 0, "max_delay": 0,
            }
            app.account_states = client_app.AccountStateStore()
            app.account_states.sync([{"stt": 1, "account_id": "raw-user"}])
            app.account_log_context = contextvars.ContextVar("test_account", default=None)
            app.global_logs = []
            app.global_log_lock = threading.RLock()
            app.selected_log_account = None
            app.post_ui = lambda _callback: None
            app.refresh_account_state_row = lambda _index: None
            app.render_log_view = lambda: None
            app.update_tree_row = lambda *_args, **_kwargs: None

            page = FakePage(final_url, login_error=login_error)
            context = FakeContext()

            async def create_browser_page(*_args, **_kwargs):
                return FakeBrowser(), context, page

            async def get_current_friends_count(_page):
                return 0

            app.create_browser_page = create_browser_page
            app.get_current_friends_count = get_current_friends_count
            await app.process_account_scoped(
                object(), 1, "raw-user", "FACEBOOK|raw-user|secret", "",
                asyncio.Semaphore(1), account_type="RAW",
                login_user="raw-user", login_password="secret",
            )
            return app.account_states.get(1)

        login_cases = [
            ("valid", "https://www.facebook.com/", None, "LIVE"),
            ("disabled", "https://www.facebook.com/disabled/", None, "DIE"),
            ("login-wall", "https://www.facebook.com/login/", None, "DIE"),
            ("checkpoint", "https://www.facebook.com/checkpoint/", None, "DIE"),
            ("challenge", "https://www.facebook.com/challenge/", None, "DIE"),
            ("timeout", "https://www.facebook.com/", TimeoutError("navigation timeout"), "ERROR"),
            ("network", "https://www.facebook.com/", OSError("DNS failure"), "ERROR"),
            ("browser", "https://www.facebook.com/", RuntimeError("browser crashed"), "ERROR"),
        ]
        for case_name, final_url, login_error, expected_status in login_cases:
            with self.subTest(case=case_name):
                state = asyncio.run(run_login_case(final_url, login_error=login_error))
                self.assertEqual(state["status"], expected_status)
                if expected_status != "LIVE":
                    self.assertNotEqual(state["current_action"], "Hoàn thành")

        class FriendButton:
            def __init__(self):
                self.clicked = False

            async def get_attribute(self, name):
                if name == "aria-label":
                    return "รอดำเนินการ" if self.clicked else "เพิ่มเป็นเพื่อน"
                if name == "aria-pressed":
                    return "true" if self.clicked else "false"
                return ""

            async def inner_text(self, **_kwargs):
                return "รอดำเนินการ" if self.clicked else "เพิ่มเป็นเพื่อน"

            async def click(self, **_kwargs):
                self.clicked = True

            async def count(self):
                return 1

            async def is_visible(self):
                return True

        class EmptyLocator:
            @property
            def last(self):
                return self

            async def count(self):
                return 0

        class FriendPage:
            def locator(self, _selector):
                return EmptyLocator()

        friend_app = client_app.MainToolApp.__new__(client_app.MainToolApp)
        with mock.patch.object(client_app.asyncio, "sleep", new=mock.AsyncMock()):
            confirmed, _detail = asyncio.run(
                friend_app.click_and_confirm_friend_request(FriendPage(), FriendButton())
            )
        self.assertTrue(confirmed)

    def test_create_page_button_languages(self):
        for label in (
            "Tạo Trang", "Create Page", "Créer une Page", "Crear página", "Criar Página",
            "สร้างเพจ", "Buat Halaman", "Gumawa ng Page", "ページを作成", "페이지 만들기",
        ):
            self.assertIsNotNone(client_app.CREATE_PAGE_BUTTON_PATTERN.fullmatch(label))

    def test_page_plan_supports_name_and_category(self):
        self.assertEqual(
            client_app.parse_page_plan("Cửa hàng A|Dịch vụ địa phương"),
            {"name": "Cửa hàng A", "category": "Dịch vụ địa phương"},
        )
        self.assertEqual(client_app.parse_page_plan("Cửa hàng B")["category"], "Blog cá nhân")
        valid, reason = client_app.validate_create_page_targets(["Cửa hàng A|Spa"], 15)
        self.assertTrue(valid, reason)
        with mock.patch.object(
            client_app,
            "generate_random_person_name",
            side_effect=["Page tự động 1", "Page tự động 2"],
        ):
            generated_plans = client_app.build_create_page_plans([], 2)
        self.assertEqual(
            generated_plans,
            ["Page tự động 1|Blog cá nhân", "Page tự động 2|Blog cá nhân"],
        )
        with mock.patch.object(
            client_app,
            "generate_random_person_name",
            return_value="Page tự động",
        ):
            filled_plans = client_app.build_create_page_plans(["Cửa hàng A|Spa"], 3)
        self.assertEqual(filled_plans[0], "Cửa hàng A|Spa")
        self.assertEqual(len(filled_plans), 3)
        mixed_targets = [
            "by_uid:10001",
            "Cửa hàng A|Spa",
            "friend_request:https://www.facebook.com/profile.php?id=10002",
        ]
        self.assertEqual(
            client_app.filter_targets_for_mode(
                mixed_targets,
                "create_page",
                {"by_uid": True, "friend_request": True, "create_page": True},
            ),
            ["Cửa hàng A|Spa"],
        )
        self.assertEqual(
            client_app.filter_targets_for_mode(
                [" CREATE_PAGE : Cửa hàng B|Nhà hàng", "Cửa hàng cũ|Blog"],
                "create_page",
                {"create_page": True},
            ),
            ["Cửa hàng B|Nhà hàng"],
        )

    def test_page_context_requires_verified_identity_for_success(self):
        context = client_app.build_page_context(
            "account-A", "Trang A", "Spa", proxy="10.0.0.1:8001", locale="th-TH"
        )
        self.assertEqual(context["creation_status"], "PENDING")
        client_app.transition_page_context(context, "VALIDATING")
        client_app.transition_page_context(context, "VERIFYING")
        with self.assertRaises(ValueError):
            client_app.transition_page_context(context, "SUCCESS", {"url": "", "id": ""})
        client_app.transition_page_context(
            context,
            "SUCCESS",
            {"url": "https://www.facebook.com/profile.php?id=123456789", "id": "123456789"},
        )
        self.assertEqual(context["owner_account_id"], "account-A")
        self.assertEqual(context["requested_category"], "Spa")
        self.assertEqual(context["locale"], "th-TH")
        self.assertEqual(context["verification_status"], "VERIFIED")

    def test_page_count_semantics_one_five_and_fifteen(self):
        with mock.patch.object(client_app, "generate_random_person_name", return_value="Auto Page"):
            for requested in (1, 5, 15):
                plans = client_app.build_create_page_plans([], requested)
                self.assertEqual(len(plans), requested)
                self.assertTrue(client_app.validate_create_page_targets(plans, requested)[0])

    def test_page_access_feedback_is_not_false_success(self):
        self.assertEqual(client_app.classify_page_access_feedback("Invitation sent"), "INVITED")
        self.assertEqual(client_app.classify_page_access_feedback("กำลังรอดำเนินการ"), "PENDING")
        self.assertEqual(client_app.classify_page_access_feedback("대기 중"), "PENDING")
        self.assertEqual(client_app.classify_page_access_feedback("Button clicked"), "FAILED")

    def test_page_access_result_keeps_owner_page_and_target(self):
        result = client_app.build_page_access_result(
            "account-A", "https://www.facebook.com/page-a", "10002", "PENDING"
        )
        self.assertEqual(result["account_id"], "account-A")
        self.assertEqual(result["page_id"], "page-a")
        self.assertEqual(result["target_user"], "10002")
        self.assertEqual(result["business_id"], "")

    def test_page_and_access_worker_limits(self):
        self.assertEqual(
            client_app.effective_account_worker_count(
                20, {"create_page": True, "add_page_admin": True}, 3, 2
            ),
            2,
        )
        self.assertEqual(
            client_app.effective_account_worker_count(
                20, {"create_page": False, "add_page_admin": True}, 3, 2
            ),
            2,
        )

    def test_page_duplicate_keys_use_verified_identity(self):
        first = client_app.page_identity_keys(
            "https://www.facebook.com/Page-A/", "123456789"
        )
        second = client_app.page_identity_keys(
            "https://www.facebook.com/page-a", "123456789"
        )
        self.assertTrue(first & second)
        self.assertFalse(client_app.page_identity_keys("", ""))

    def test_page_url_without_numeric_id_is_verified(self):
        identity = client_app.extract_facebook_page_identity(
            ["https://www.facebook.com/my-valid-page"]
        )
        self.assertEqual(identity["url"], "https://www.facebook.com/my-valid-page")
        self.assertTrue(client_app.is_verified_page_identity(identity))

    def test_page_missing_identity_cannot_be_success(self):
        context = client_app.build_page_context("account-A", "Page A", "Spa")
        with self.assertRaises(ValueError):
            client_app.transition_page_context(context, "SUCCESS", None)
        self.assertNotEqual(context["creation_status"], "SUCCESS")

    def test_create_page_transient_retry_uses_bounded_backoff(self):
        calls = []

        async def operation():
            calls.append(len(calls) + 1)
            if len(calls) < 3:
                raise TimeoutError("temporary navigation timeout")
            return "ok"

        with mock.patch.object(client_app.asyncio, "sleep", new=mock.AsyncMock()) as sleep:
            result, retry_count = asyncio.run(
                client_app.retry_create_page_operation(operation, max_attempts=3, base_delay=1)
            )
        self.assertEqual(result, "ok")
        self.assertEqual(retry_count, 2)
        self.assertEqual(calls, [1, 2, 3])
        self.assertEqual([call.args[0] for call in sleep.await_args_list], [1, 2])

    def test_create_page_permanent_error_is_not_retried(self):
        calls = []

        async def operation():
            calls.append(1)
            raise ValueError("invalid page name")

        with mock.patch.object(client_app.asyncio, "sleep", new=mock.AsyncMock()) as sleep:
            with self.assertRaises(ValueError):
                asyncio.run(client_app.retry_create_page_operation(operation, max_attempts=3))
        self.assertEqual(len(calls), 1)
        sleep.assert_not_awaited()

    def test_create_page_network_errors_are_technical(self):
        self.assertTrue(client_app.is_transient_create_page_error(TimeoutError("timeout")))
        self.assertTrue(client_app.is_transient_create_page_error(OSError("DNS failure")))
        self.assertFalse(client_app.is_transient_create_page_error(ValueError("invalid category")))
        self.assertEqual(client_app.account_status_for_failure("proxy"), "ERROR")

    def test_parallel_page_contexts_do_not_cross_account(self):
        context_a = client_app.build_page_context("A", "Page A", "Spa", proxy="proxy-A")
        context_b = client_app.build_page_context("B", "Page B", "Restaurant", proxy="proxy-B")
        client_app.transition_page_context(
            context_a, "SUCCESS", {"url": "https://www.facebook.com/page-a", "id": "11111"}
        )
        self.assertEqual(context_a["page_id"], "11111")
        self.assertEqual(context_b["page_id"], "")
        self.assertEqual(context_b["proxy"], "proxy-B")
        self.assertEqual(context_b["requested_category"], "Restaurant")

    def test_cancelled_page_context_is_never_verified(self):
        context = client_app.build_page_context("A", "Page A", "Spa")
        client_app.transition_page_context(context, "CANCELLED")
        self.assertEqual(context["creation_status"], "CANCELLED")
        self.assertEqual(context["verification_status"], "NOT_VERIFIED")

    def test_page_admin_mapping_rejects_other_owner(self):
        jobs = ["1|https://www.facebook.com/page-a|10001"]
        self.assertIsNone(client_app.select_page_admin_job(jobs, 2, "account-B"))
        self.assertFalse(client_app.page_reference_matches(
            "https://www.facebook.com/page-a",
            "https://www.facebook.com/page-b/settings/?tab=profile_access",
        ))

    def test_page_admin_job_is_scoped_to_account(self):
        targets = [
            "1|https://www.facebook.com/page-one|10001",
            "2|AUTO|10002",
        ]
        self.assertEqual(
            client_app.select_page_admin_job(targets, 2, "FB_2"),
            {"page": "AUTO", "admin": "10002"},
        )
        self.assertIsNone(client_app.select_page_admin_job(targets, 3, "FB_3"))
        self.assertTrue(client_app.page_reference_matches(
            "https://www.facebook.com/page-one",
            "https://www.facebook.com/page-one/settings/?tab=profile_access",
        ))
        self.assertFalse(client_app.page_reference_matches(
            "https://www.facebook.com/page-one",
            "https://www.facebook.com/page-two/settings/?tab=profile_access",
        ))

        class TargetLocator:
            async def get_attribute(self, name):
                return "/profile.php?id=10002" if name == "href" else ""

            async def inner_text(self, **_kwargs):
                return "บัญชีเป้าหมาย"

        self.assertTrue(asyncio.run(
            client_app.locator_matches_account_target(TargetLocator(), "10002")
        ))
        self.assertFalse(asyncio.run(
            client_app.locator_matches_account_target(TargetLocator(), "10001")
        ))

    def test_page_identity_rejects_creation_screen(self):
        self.assertEqual(
            client_app.extract_facebook_page_identity(
                ["https://www.facebook.com/pages/creation/"]
            ),
            {"url": "", "id": ""},
        )
        self.assertEqual(
            client_app.extract_facebook_page_identity(
                ["https://www.facebook.com/profile.php?id=123456789"]
            ),
            {
                "url": "https://www.facebook.com/profile.php?id=123456789",
                "id": "123456789",
            },
        )
        self.assertFalse(client_app.is_verified_page_identity({"url": "", "id": ""}))
        self.assertTrue(
            client_app.is_verified_page_identity(
                {"url": "https://www.facebook.com/profile.php?id=123456789", "id": "123456789"}
            )
        )

        class FakeKeyboard:
            def __init__(self, page):
                self.page = page

            async def press(self, *_args, **_kwargs):
                return None

            async def type(self, value, **_kwargs):
                if self.page.active_field == "category":
                    self.page.category_buffer += value
                return None

        class FakeLocator:
            def __init__(
                self, count=0, on_click=None, on_focus=None,
                on_fill=None, click_error=None,
            ):
                self._count = count
                self._on_click = on_click
                self._on_focus = on_focus
                self._on_fill = on_fill
                self._click_error = click_error

            @property
            def first(self):
                return self

            def nth(self, _index):
                return self

            async def count(self):
                return self._count

            async def is_visible(self):
                return self._count > 0

            async def wait_for(self, **_kwargs):
                if self._count == 0:
                    raise TimeoutError("not visible")

            async def scroll_into_view_if_needed(self):
                return None

            async def click(self, **_kwargs):
                if self._click_error:
                    raise self._click_error
                if self._on_click:
                    self._on_click()

            async def focus(self):
                if self._on_focus:
                    self._on_focus()
                return None

            async def fill(self, value):
                if self._on_fill:
                    self._on_fill(value)
                return None

            async def get_attribute(self, _name):
                return None

        class FakePage:
            def __init__(self, result_url=None, body_text="", click_error=None):
                self.url = "https://www.facebook.com/pages/creation/"
                self.frames = []
                self.active_field = ""
                self.category_buffer = ""
                self.typed_categories = []
                self.keyboard = FakeKeyboard(self)
                self.result_urls = result_url if isinstance(result_url, list) else [result_url]
                self.create_index = 0
                self.body_text = body_text
                self.click_error = click_error

            async def goto(self, *_args, **_kwargs):
                self.url = "https://www.facebook.com/pages/creation/"

            async def title(self):
                return "Create a Page"

            async def inner_text(self, _selector):
                return self.body_text

            def locator(self, selector):
                if "Tên trang" in selector or "Page name" in selector:
                    return FakeLocator(
                        count=1,
                        on_focus=lambda: setattr(self, "active_field", "name"),
                    )
                if "Hạng mục" in selector or "Category" in selector:
                    def focus_category():
                        self.active_field = "category"
                        self.category_buffer = ""

                    return FakeLocator(
                        count=1,
                        on_focus=focus_category,
                        on_fill=lambda value: setattr(self, "category_buffer", value),
                    )
                if 'role="listbox"' in selector:
                    return FakeLocator(
                        count=1,
                        on_click=lambda: self.typed_categories.append(self.category_buffer),
                    )
                if selector == "input":
                    return FakeLocator(count=1)
                return FakeLocator(count=0)

            def get_by_role(self, role, **_kwargs):
                if role != "button":
                    return FakeLocator(count=0)

                def finish_create():
                    result_url = self.result_urls[self.create_index]
                    self.create_index += 1
                    if result_url:
                        self.url = result_url

                return FakeLocator(
                    count=1,
                    on_click=finish_create,
                    click_error=self.click_error,
                )

        async def run_create_case(
            result_url=None,
            body_text="",
            click_error=None,
            targets=None,
            running=True,
        ):
            app = client_app.MainToolApp.__new__(client_app.MainToolApp)
            app.is_running = running
            app.run_config = {
                "feed_surf_min": 0,
                "watch_review_min": 0,
                "modes": {"create_page": True},
            }
            app.log = mock.Mock()
            app.account_states = client_app.AccountStateStore()
            app.account_states.sync([
                {
                    "stt": 1,
                    "account_id": "FB_1",
                    "raw_line": "FB_1|password",
                    "source_line": "FB_1|password",
                }
            ])
            app.create_page_result_lock = threading.RLock()
            app.create_page_account_results = {"completed": {}, "die": {}}
            app.post_ui = lambda _callback: None
            app.refresh_account_state_row = lambda _index: None
            app.render_log_view = lambda: None
            app.selected_log_account = None

            async def take_error_snapshot(*_args, **_kwargs):
                return None

            app.take_error_snapshot = take_error_snapshot

            async def ensure_personal_profile(*_args, **_kwargs):
                return True

            app.ensure_personal_profile = ensure_personal_profile
            fake_page = FakePage(result_url, body_text, click_error)
            page_targets = targets or ["Trang kiểm thử"]
            with (
                mock.patch.object(client_app.asyncio, "sleep", new=mock.AsyncMock()),
                mock.patch.object(
                    client_app, "save_created_page_success", return_value=True
                ) as save_result,
                mock.patch.object(client_app, "save_create_page_outcome"),
                mock.patch.object(client_app, "append_create_page_account_log"),
            ):
                records = await app.run_create_page(
                    fake_page,
                    object(),
                    "FB_1",
                    1,
                    targets=page_targets,
                    max_pages=len(page_targets),
                    min_page_del=0,
                    max_page_del=0,
                )
            return (
                records,
                save_result.call_count,
                fake_page.typed_categories,
                app.account_states.get(1)["status"],
            )

        cases = (
            ("manage text only", None, "manage page / quản lý trang", None, "FAILED", 0, "UNKNOWN"),
            (
                "verified URL",
                "https://www.facebook.com/profile.php?id=123456789",
                "",
                None,
                "SUCCESS",
                1,
                "UNKNOWN",
            ),
            ("create exception", None, "", RuntimeError("click failed"), "ERROR", 0, "UNKNOWN"),
            (
                "account disabled during create",
                "https://www.facebook.com/disabled/",
                "",
                None,
                "FAILED",
                0,
                "DIE",
            ),
        )
        for (
            case_name, result_url, body_text, click_error,
            expected_status, expected_saves, expected_account_status,
        ) in cases:
            with self.subTest(case=case_name):
                records, save_calls, _categories, account_status = asyncio.run(
                    run_create_case(result_url, body_text, click_error)
                )
                self.assertEqual(len(records), 1)
                self.assertEqual(records[0]["status"], expected_status)
                self.assertEqual(save_calls, expected_saves)
                self.assertEqual(account_status, expected_account_status)

        category_cases = (
            ("restaurant", "Trang A|Nhà hàng", "Nhà hàng", "123456789"),
            ("spa", "Trang B|Spa", "Spa", "223456789"),
            ("empty fallback", "Trang C|", "Blog cá nhân", "323456789"),
        )
        for case_name, target, expected_category, page_id in category_cases:
            with self.subTest(category=case_name):
                records, _save_calls, categories, _account_status = asyncio.run(
                    run_create_case(
                        f"https://www.facebook.com/profile.php?id={page_id}",
                        targets=[target],
                    )
                )
                self.assertEqual(categories, [expected_category])
                self.assertEqual(records[0]["category"], expected_category)
                if expected_category != "Blog":
                    self.assertNotEqual(categories, ["Blog"])

        records, _save_calls, categories, _account_status = asyncio.run(
            run_create_case(
                [
                    "https://www.facebook.com/profile.php?id=423456789",
                    "https://www.facebook.com/profile.php?id=523456789",
                ],
                targets=["Trang A|Nhà hàng", "Trang B|Spa"],
            )
        )
        self.assertEqual(categories, ["Nhà hàng", "Spa"])
        self.assertEqual([record["category"] for record in records], ["Nhà hàng", "Spa"])

        cancelled_records, cancelled_saves, _categories, _status = asyncio.run(
            run_create_case(
                "https://www.facebook.com/profile.php?id=623456789",
                running=False,
            )
        )
        self.assertEqual(cancelled_records, [])
        self.assertEqual(cancelled_saves, 0)

        self.assertEqual(
            client_app.validate_create_page_targets(["Trang A|Spa"], 1),
            (True, ""),
        )
        self.assertEqual(
            client_app.validate_create_page_targets(["Trang A|Spa"], 5),
            (True, ""),
        )
        self.assertFalse(client_app.validate_create_page_targets(["|Spa"], 1)[0])

        success_record = client_app.build_create_page_result(
            "SUCCESS", "FB_1", "Trang A", "Spa",
            page_url="https://www.facebook.com/profile.php?id=623456789",
            page_id="623456789",
        )
        with tempfile.TemporaryDirectory() as temp_dir, mock.patch.object(
            client_app, "APP_DATA_DIR", temp_dir
        ):
            self.assertTrue(client_app.save_created_page_success(success_record))
            self.assertFalse(client_app.save_created_page_success(success_record))

        attempts = []

        async def transient_operation():
            attempts.append(len(attempts) + 1)
            if len(attempts) < 3:
                raise TimeoutError("temporary navigation timeout")
            return "ok"

        with mock.patch.object(client_app.asyncio, "sleep", new=mock.AsyncMock()):
            value, retry_count = asyncio.run(client_app.retry_create_page_operation(
                transient_operation, max_attempts=3, base_delay=0
            ))
        self.assertEqual((value, retry_count), ("ok", 2))

        class Closable:
            def __init__(self):
                self.closed = False

            async def close(self):
                self.closed = True

        fake_context, fake_browser = Closable(), Closable()
        asyncio.run(client_app.close_browser_resources(fake_context, fake_browser))
        self.assertTrue(fake_context.closed)
        self.assertTrue(fake_browser.closed)

    def test_async_methods_do_not_read_tk_widgets(self):
        source_path = pathlib.Path(client_app.__file__)
        tree = ast.parse(source_path.read_text(encoding="utf-8"))
        forbidden_prefixes = (
            "self.ent_",
            "self.txt_",
            "self.chk_",
            "self.proxy_mode",
            "self.mode_vars",
            "self.tree",
            "self.root",
        )
        violations = []
        for function in ast.walk(tree):
            if not isinstance(function, ast.AsyncFunctionDef):
                continue
            for node in ast.walk(function):
                if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Attribute):
                    continue
                owner = ast.unparse(node.func.value)
                if node.func.attr == "get" and owner.startswith(forbidden_prefixes):
                    violations.append(f"{function.name}:{node.lineno}:{owner}.get")
        self.assertEqual(violations, [])


if __name__ == "__main__":
    unittest.main()
